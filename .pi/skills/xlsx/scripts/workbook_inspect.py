"""Inspect an OOXML workbook for preservation risks, structure, and formula-quality issues."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

from office.helpers.xlsx_package import (
    EXCEL_PACKAGE_EXTENSIONS,
    content_type_overrides,
    is_probably_encrypted_excel,
    is_zip_package,
    list_package_parts,
    package_feature_counts,
    sheet_part_map,
)

ERROR_LITERALS = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#CONNECT!",
    "#UNKNOWN!",
    "#BUSY!",
    "#GETTING_DATA",
}
VOLATILE_FUNCTIONS = {
    "AREAS",
    "CELL",
    "INFO",
    "INDIRECT",
    "NOW",
    "OFFSET",
    "RAND",
    "RANDBETWEEN",
    "TODAY",
}
EXTERNAL_BOOK_RE = re.compile(r"\[[^\]]+\][^!]+!")


def _iso(ts: float | None) -> str | None:
    if ts is None:
        return None
    return datetime.fromtimestamp(ts).isoformat(timespec="seconds")


def _cell_values(ws):
    cells = getattr(ws, "_cells", None)
    if cells:
        return list(cells.values())
    return []


def _table_records(ws) -> list[dict[str, Any]]:
    tables = []
    for table in ws.tables.values():
        style = getattr(getattr(table, "tableStyleInfo", None), "name", None)
        tables.append(
            {
                "name": table.displayName,
                "ref": table.ref,
                "style": style,
                "totals_row": bool(getattr(table, "totalsRowShown", False)),
            }
        )
    return tables


def _detect_overlapping_ranges(ranges: list[str]) -> list[list[str]]:
    overlaps: list[list[str]] = []
    parsed = []
    for ref in ranges:
        try:
            parsed.append((ref, range_boundaries(ref)))
        except Exception:
            continue
    for i, (left_ref, (a1, b1, c1, d1)) in enumerate(parsed):
        for right_ref, (a2, b2, c2, d2) in parsed[i + 1 :]:
            if not (c1 < a2 or c2 < a1 or d1 < b2 or d2 < b1):
                overlaps.append([left_ref, right_ref])
    return overlaps


def _defined_name_records(wb) -> list[dict[str, Any]]:
    records = []
    for name, defn in wb.defined_names.items():
        records.append(
            {
                "name": name,
                "text": getattr(defn, "attr_text", None),
                "local_sheet_id": getattr(defn, "localSheetId", None),
                "hidden": getattr(defn, "hidden", None),
            }
        )
    return records


def _formula_lints_for_cell(formula: str) -> dict[str, Any]:
    volatile: list[str] = []
    suspicious_numbers: list[str] = []
    try:
        tokens = Tokenizer(formula).items
        for idx, token in enumerate(tokens):
            if token.type == "FUNC" and token.subtype == "OPEN":
                func_name = token.value[:-1].upper()
                if func_name in VOLATILE_FUNCTIONS:
                    volatile.append(func_name)
            elif token.type == "OPERAND" and token.subtype == "NUMBER":
                try:
                    numeric = float(token.value)
                except Exception:
                    continue
                if numeric not in {0.0, 1.0, -1.0}:
                    suspicious_numbers.append(token.value)
    except Exception:
        pass

    return {
        "volatile_functions": sorted(set(volatile)),
        "has_external_reference": bool(EXTERNAL_BOOK_RE.search(formula)),
        "formula_error_literals": sorted(err for err in ERROR_LITERALS if err in formula),
        "suspicious_numeric_literals": suspicious_numbers,
    }


def _workbook_security_info(wb) -> dict[str, Any]:
    sec = wb.security
    return {
        "lock_structure": bool(getattr(sec, "lockStructure", False) or getattr(sec, "lock_structure", False)),
        "lock_windows": bool(getattr(sec, "lockWindows", False) or getattr(sec, "lock_windows", False)),
        "lock_revision": bool(getattr(sec, "lockRevision", False) or getattr(sec, "lock_revision", False)),
        "has_workbook_password_hash": bool(getattr(sec, "workbookPassword", None) or getattr(sec, "workbookHashValue", None)),
    }


def _inspect_workbook_internal(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "file": {
            "path": str(path.resolve()),
            "name": path.name,
            "suffix": path.suffix.lower(),
            "size_bytes": path.stat().st_size if path.exists() else None,
            "modified_at": _iso(path.stat().st_mtime) if path.exists() else None,
        },
        "status": "success",
        "issues": [],
        "warnings": [],
        "package": {},
        "workbook": {},
        "sheets": [],
        "formula_lint_summary": {},
        "comparison": None,
    }

    if not path.exists():
        result["status"] = "error"
        result["issues"].append(f"Workbook does not exist: {path}")
        return result

    if path.suffix.lower() not in EXCEL_PACKAGE_EXTENSIONS:
        result["status"] = "error"
        result["issues"].append(
            f"Unsupported file type '{path.suffix}'. Expected one of: {', '.join(sorted(EXCEL_PACKAGE_EXTENSIONS))}"
        )
        return result

    encrypted = is_probably_encrypted_excel(path)
    zipped = is_zip_package(path)
    result["package"]["is_zip_package"] = zipped
    result["package"]["is_probably_encrypted"] = encrypted

    if encrypted:
        result["status"] = "error"
        result["issues"].append(
            "Workbook appears to be encrypted/password-protected OOXML stored in an OLE container. Decrypt before editing."
        )
        return result

    if not zipped:
        result["status"] = "error"
        result["issues"].append("Workbook is not a valid OOXML ZIP package.")
        return result

    parts = list_package_parts(path)
    feature_counts = package_feature_counts(parts)
    package_sheet_map = sheet_part_map(path)
    result["package"].update(
        {
            "parts": parts,
            "feature_counts": feature_counts,
            "sheet_parts": package_sheet_map,
            "content_type_overrides": content_type_overrides(path),
        }
    )

    if feature_counts.get("has_vba"):
        result["warnings"].append("Workbook contains VBA. Use openpyxl with keep_vba=True and preserve the original extension.")
    if feature_counts.get("pivot_tables") or feature_counts.get("pivot_cache_defs"):
        result["warnings"].append(
            "Workbook contains pivot tables/caches. openpyxl can generally preserve them, but creating or rebuilding pivots is not a safe editing path."
        )
    if feature_counts.get("external_links") or feature_counts.get("connections") or feature_counts.get("query_tables"):
        result["warnings"].append(
            "Workbook contains external links / connections / query tables. Avoid destructive sheet or range surgery without a before/after compare."
        )
    if feature_counts.get("charts") or feature_counts.get("images") or feature_counts.get("drawings"):
        result["warnings"].append(
            "Workbook contains charts/images/drawings. Prefer surgical edits and review a rendered preview afterward."
        )

    keep_vba = path.suffix.lower() in {".xlsm", ".xltm"}
    wb = load_workbook(path, data_only=False, keep_links=True, keep_vba=keep_vba)
    try:
        calc = wb.calculation
        result["workbook"] = {
            "sheet_order": list(wb.sheetnames),
            "sheet_count": len(wb.sheetnames),
            "epoch": getattr(wb.epoch, "isoformat", lambda: str(wb.epoch))(),
            "template": bool(getattr(wb, "template", False)),
            "style_count": len(getattr(wb, "_cell_styles", [])),
            "named_style_count": len(getattr(wb, "_named_styles", [])),
            "defined_names": _defined_name_records(wb),
            "calc_properties": {
                "calc_id": getattr(calc, "calcId", None),
                "calc_mode": getattr(calc, "calcMode", None),
                "calc_on_save": getattr(calc, "calcOnSave", None),
                "full_calc_on_load": getattr(calc, "fullCalcOnLoad", None),
                "force_full_calc": getattr(calc, "forceFullCalc", None),
                "iterate": getattr(calc, "iterate", None),
            },
            "security": _workbook_security_info(wb),
        }

        if len(package_sheet_map) != len(wb.sheetnames):
            result["warnings"].append(
                f"Workbook package lists {len(package_sheet_map)} sheets but openpyxl exposed {len(wb.sheetnames)} names. Review chartsheets and specialized sheet types carefully."
            )

        if getattr(calc, "calcMode", None) == "manual":
            result["warnings"].append("Workbook calculation mode is manual. Force a recalc and cache values before delivery.")

        sheet_formula_total = 0
        workbook_cell_errors = Counter()
        volatile_counter = Counter()
        suspicious_literal_counter = Counter()
        external_formula_locations: list[str] = []
        literal_error_formula_locations: list[str] = []
        duplicate_table_names: list[str] = []
        seen_table_names: set[str] = set()

        for ws in wb.worksheets:
            cells = _cell_values(ws)
            tables = _table_records(ws)
            merged_refs = [str(rng) for rng in ws.merged_cells.ranges]
            merged_overlaps = _detect_overlapping_ranges(merged_refs)
            hidden_rows = sum(1 for dim in ws.row_dimensions.values() if getattr(dim, "hidden", False))
            hidden_cols = sum(1 for dim in ws.column_dimensions.values() if getattr(dim, "hidden", False))
            data_validation_count = len(getattr(ws.data_validations, "dataValidation", []))
            conditional_blocks = len(ws.conditional_formatting)
            formula_cells = 0
            cell_error_locations: list[str] = []
            volatile_locations: dict[str, list[str]] = {name: [] for name in VOLATILE_FUNCTIONS}
            suspicious_literal_cells: list[dict[str, Any]] = []
            external_ref_cells: list[str] = []
            formula_error_literal_cells: list[str] = []
            comments_count = 0
            hyperlink_count = 0
            chart_count = len(getattr(ws, "_charts", []))
            image_count = len(getattr(ws, "_images", []))

            for cell in cells:
                if cell.comment is not None:
                    comments_count += 1
                if getattr(cell, "hyperlink", None) is not None:
                    hyperlink_count += 1
                if isinstance(cell.value, str) and cell.value in ERROR_LITERALS:
                    workbook_cell_errors[cell.value] += 1
                    if len(cell_error_locations) < 25:
                        cell_error_locations.append(f"{ws.title}!{cell.coordinate}")
                if cell.data_type != "f":
                    continue
                formula_cells += 1
                sheet_formula_total += 1
                formula = cell.value or ""
                lints = _formula_lints_for_cell(formula)
                if lints["has_external_reference"]:
                    external_formula_locations.append(f"{ws.title}!{cell.coordinate}")
                    if len(external_ref_cells) < 25:
                        external_ref_cells.append(f"{ws.title}!{cell.coordinate}")
                if lints["formula_error_literals"]:
                    literal_error_formula_locations.append(f"{ws.title}!{cell.coordinate}")
                    if len(formula_error_literal_cells) < 25:
                        formula_error_literal_cells.append(f"{ws.title}!{cell.coordinate}")
                for func_name in lints["volatile_functions"]:
                    volatile_counter[func_name] += 1
                    if len(volatile_locations[func_name]) < 10:
                        volatile_locations[func_name].append(f"{ws.title}!{cell.coordinate}")
                if lints["suspicious_numeric_literals"]:
                    suspicious_literal_counter.update(lints["suspicious_numeric_literals"])
                    if len(suspicious_literal_cells) < 25:
                        suspicious_literal_cells.append(
                            {
                                "cell": f"{ws.title}!{cell.coordinate}",
                                "formula": formula,
                                "numbers": lints["suspicious_numeric_literals"][:10],
                            }
                        )

            for table in tables:
                if table["name"] in seen_table_names:
                    duplicate_table_names.append(table["name"])
                else:
                    seen_table_names.add(table["name"])

            auto_filter_ref = getattr(getattr(ws, "auto_filter", None), "ref", None)
            if auto_filter_ref and any(table["ref"] == auto_filter_ref for table in tables):
                result["issues"].append(
                    f"{ws.title}: worksheet auto_filter.ref duplicates a table range ({auto_filter_ref}). Remove the worksheet-level auto filter."
                )

            if merged_overlaps:
                result["warnings"].append(
                    f"{ws.title}: overlapping merged ranges detected: {merged_overlaps[:5]}"
                )

            sheet_record = {
                "name": ws.title,
                "state": ws.sheet_state,
                "dimensions": ws.calculate_dimension(),
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "auto_filter_ref": auto_filter_ref,
                "tables": tables,
                "merged_ranges": merged_refs,
                "merged_count": len(merged_refs),
                "data_validation_count": data_validation_count,
                "conditional_format_blocks": conditional_blocks,
                "chart_count": chart_count,
                "image_count": image_count,
                "comments_count": comments_count,
                "hyperlink_count": hyperlink_count,
                "formula_cells": formula_cells,
                "volatile_formula_locations": {k: v for k, v in volatile_locations.items() if v},
                "external_reference_formula_cells": external_ref_cells,
                "formula_error_literal_cells": formula_error_literal_cells,
                "suspicious_numeric_literal_cells": suspicious_literal_cells,
                "cell_error_locations": cell_error_locations,
                "hidden_rows": hidden_rows,
                "hidden_columns": hidden_cols,
                "protection": {
                    "sheet": bool(getattr(ws.protection, "sheet", False)),
                    "objects": bool(getattr(ws.protection, "objects", False)),
                    "scenarios": bool(getattr(ws.protection, "scenarios", False)),
                },
                "print_settings": {
                    "orientation": getattr(ws.page_setup, "orientation", None),
                    "paper_size": getattr(ws.page_setup, "paperSize", None),
                    "fit_to_width": getattr(ws.page_setup, "fitToWidth", None),
                    "fit_to_height": getattr(ws.page_setup, "fitToHeight", None),
                    "print_title_rows": ws.print_title_rows,
                    "print_title_cols": ws.print_title_cols,
                },
            }
            result["sheets"].append(sheet_record)

        for defn in result["workbook"]["defined_names"]:
            text = (defn.get("text") or "")
            if "#REF!" in text:
                result["issues"].append(f"Defined name '{defn['name']}' contains #REF!: {text}")
            for sheet_name in re.findall(r"(?:'((?:[^']|'')+)'|([^'!,\[]+))!", text):
                candidate = (sheet_name[0] or sheet_name[1] or "").replace("''", "'")
                if candidate and candidate not in wb.sheetnames:
                    result["issues"].append(
                        f"Defined name '{defn['name']}' references missing sheet '{candidate}'"
                    )

        if duplicate_table_names:
            unique_dupes = sorted(set(duplicate_table_names))
            result["issues"].append(f"Duplicate table names detected: {unique_dupes}")

        if workbook_cell_errors:
            result["issues"].append(
                f"Workbook contains cached cell errors: {dict(workbook_cell_errors)}"
            )

        result["formula_lint_summary"] = {
            "total_formula_cells": sheet_formula_total,
            "volatile_functions": dict(volatile_counter),
            "external_reference_formula_cells": len(external_formula_locations),
            "formula_error_literal_cells": len(literal_error_formula_locations),
            "suspicious_numeric_literals": dict(suspicious_literal_counter.most_common(25)),
            "cached_cell_errors": dict(workbook_cell_errors),
            "sample_external_reference_cells": external_formula_locations[:25],
            "sample_formula_error_literal_cells": literal_error_formula_locations[:25],
        }

    finally:
        wb.close()

    if result["issues"]:
        result["status"] = "issues_found"
    return result


def _load_baseline(baseline_path: Path) -> dict[str, Any]:
    if baseline_path.suffix.lower() == ".json":
        return json.loads(baseline_path.read_text(encoding="utf-8"))
    return _inspect_workbook_internal(baseline_path)


def _sheet_summary_map(inspection: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {sheet["name"]: sheet for sheet in inspection.get("sheets", [])}


def compare_inspections(current: dict[str, Any], baseline: dict[str, Any]) -> dict[str, Any]:
    current_parts = set(current.get("package", {}).get("parts", []))
    baseline_parts = set(baseline.get("package", {}).get("parts", []))
    current_features = current.get("package", {}).get("feature_counts", {})
    baseline_features = baseline.get("package", {}).get("feature_counts", {})
    current_workbook = current.get("workbook", {})
    baseline_workbook = baseline.get("workbook", {})
    current_sheets = _sheet_summary_map(current)
    baseline_sheets = _sheet_summary_map(baseline)

    sheet_names = sorted(set(current_sheets) | set(baseline_sheets))
    per_sheet: dict[str, Any] = {}
    for name in sheet_names:
        before = baseline_sheets.get(name)
        after = current_sheets.get(name)
        if before is None:
            per_sheet[name] = {"status": "added", "after": after}
            continue
        if after is None:
            per_sheet[name] = {"status": "removed", "before": before}
            continue
        delta = {
            "status": "existing",
            "state_changed": before.get("state") != after.get("state"),
            "dimension_before": before.get("dimensions"),
            "dimension_after": after.get("dimensions"),
            "formula_cells_before": before.get("formula_cells"),
            "formula_cells_after": after.get("formula_cells"),
            "tables_before": sorted(table["name"] for table in before.get("tables", [])),
            "tables_after": sorted(table["name"] for table in after.get("tables", [])),
            "chart_count_before": before.get("chart_count"),
            "chart_count_after": after.get("chart_count"),
            "image_count_before": before.get("image_count"),
            "image_count_after": after.get("image_count"),
            "comments_count_before": before.get("comments_count"),
            "comments_count_after": after.get("comments_count"),
            "hyperlink_count_before": before.get("hyperlink_count"),
            "hyperlink_count_after": after.get("hyperlink_count"),
            "data_validation_count_before": before.get("data_validation_count"),
            "data_validation_count_after": after.get("data_validation_count"),
            "conditional_format_blocks_before": before.get("conditional_format_blocks"),
            "conditional_format_blocks_after": after.get("conditional_format_blocks"),
            "merged_count_before": before.get("merged_count"),
            "merged_count_after": after.get("merged_count"),
        }
        per_sheet[name] = delta

    feature_changes = {
        key: {"before": baseline_features.get(key), "after": current_features.get(key)}
        for key in sorted(set(current_features) | set(baseline_features))
        if baseline_features.get(key) != current_features.get(key)
    }

    return {
        "added_parts": sorted(current_parts - baseline_parts),
        "removed_parts": sorted(baseline_parts - current_parts),
        "sheet_order_before": baseline_workbook.get("sheet_order", []),
        "sheet_order_after": current_workbook.get("sheet_order", []),
        "sheet_order_changed": baseline_workbook.get("sheet_order", []) != current_workbook.get("sheet_order", []),
        "feature_changes": feature_changes,
        "per_sheet": per_sheet,
    }


def inspect_workbook(path: str | Path, compare_to: str | Path | None = None) -> dict[str, Any]:
    workbook_path = Path(path)
    result = _inspect_workbook_internal(workbook_path)
    if compare_to:
        baseline_path = Path(compare_to)
        if baseline_path.exists():
            result["comparison"] = compare_inspections(result, _load_baseline(baseline_path))
        else:
            result["warnings"].append(f"Baseline comparison target does not exist: {baseline_path}")
    if result["issues"] and result["status"] != "error":
        result["status"] = "issues_found"
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect an Excel OOXML workbook for structure and quality risks")
    parser.add_argument("workbook", help="Path to .xlsx/.xlsm/.xltx/.xltm workbook")
    parser.add_argument(
        "--compare-to",
        help="Optional baseline workbook or prior inspect JSON to compare against",
    )
    args = parser.parse_args()
    print(json.dumps(inspect_workbook(args.workbook, args.compare_to), indent=2))


if __name__ == "__main__":
    main()
