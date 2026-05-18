"""Recalculate spreadsheet formulas and cache results back into the original OOXML workbook."""

from __future__ import annotations

import argparse
import json
import math
import os
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import lxml.etree as ET
from openpyxl import load_workbook
from openpyxl.utils.datetime import time_to_days, to_excel

from office.helpers.xlsx_package import (
    EXCEL_PACKAGE_EXTENSIONS,
    NS,
    OOXML_MAIN_NS,
    is_probably_encrypted_excel,
    is_zip_package,
    sheet_part_map,
)
from office.soffice import initialize_profile, run_soffice, temporary_profile

ERROR_LITERALS = {
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#CONNECT!", "#UNKNOWN!",
    "#BUSY!", "#GETTING_DATA",
}

@dataclass
class CachedValue:
    xml_type: str | None
    xml_value: str | None
    python_value: Any

class RecalcError(RuntimeError):
    pass


def _format_number(value: int | float | Decimal) -> str:
    if isinstance(value, Decimal):
        return format(value, 'f')
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            raise ValueError(f'Invalid numeric result {value!r}')
        return format(value, '.15g')
    raise TypeError(value)


def _coerce_cached_value(value: Any, epoch: datetime) -> CachedValue:
    if value is None:
        return CachedValue(None, None, None)
    if isinstance(value, bool):
        return CachedValue('b', '1' if value else '0', value)
    if isinstance(value, str):
        if value in ERROR_LITERALS:
            return CachedValue('e', value, value)
        return CachedValue('str', value, value)
    if isinstance(value, datetime):
        return CachedValue(None, _format_number(to_excel(value, epoch=epoch)), value)
    if isinstance(value, date):
        return CachedValue(None, _format_number(to_excel(datetime.combine(value, time.min), epoch=epoch)), value)
    if isinstance(value, time):
        return CachedValue(None, _format_number(time_to_days(value)), value)
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return CachedValue(None, _format_number(value), value)
    return CachedValue('str', str(value), value)


def _run_libreoffice_snapshot(src: Path, timeout: int) -> Path:
    with temporary_profile() as profile:
        initialize_profile(profile, timeout=min(timeout, 20))
        with tempfile.TemporaryDirectory(prefix='xlsx-recalc-') as td:
            temp_dir = Path(td)
            source_copy = temp_dir / f'input{src.suffix.lower()}'
            shutil.copy2(src, source_copy)
            common_args = ['--headless', '--norestore', '--nodefault', '--nolockcheck', '--nofirststartwizard']
            to_ods = run_soffice(
                [*common_args, '--convert-to', 'ods', '--outdir', str(temp_dir), str(source_copy)],
                profile_dir=profile, timeout=timeout, capture_output=True, text=True, check=False,
            )
            ods_path = temp_dir / 'input.ods'
            if to_ods.returncode != 0 or not ods_path.exists():
                raise RecalcError('LibreOffice failed to build an evaluated ODS snapshot: ' + (to_ods.stderr or to_ods.stdout or 'unknown error'))
            to_xlsx = run_soffice(
                [*common_args, '--convert-to', 'xlsx', '--outdir', str(temp_dir), str(ods_path)],
                profile_dir=profile, timeout=timeout, capture_output=True, text=True, check=False,
            )
            evaluated_path = temp_dir / 'input.xlsx'
            if to_xlsx.returncode != 0 or not evaluated_path.exists():
                raise RecalcError('LibreOffice failed to convert the evaluated ODS snapshot back to XLSX: ' + (to_xlsx.stderr or to_xlsx.stdout or 'unknown error'))
            fd, temp_name = tempfile.mkstemp(prefix='evaluated-snapshot-', suffix='.xlsx')
            os.close(fd)
            snapshot_copy = Path(temp_name)
            shutil.copy2(evaluated_path, snapshot_copy)
            return snapshot_copy


def _formula_cell_map(original_path: Path, evaluated_path: Path) -> tuple[dict[str, dict[str, CachedValue]], int, list[str]]:
    original_wb = load_workbook(original_path, data_only=False, keep_links=True)
    evaluated_wb = load_workbook(evaluated_path, data_only=True, keep_links=False)
    value_map: dict[str, dict[str, CachedValue]] = {}
    total_formulas = 0
    warnings: list[str] = []
    try:
        for ws in original_wb.worksheets:
            if ws.title not in evaluated_wb.sheetnames:
                warnings.append(f"Sheet '{ws.title}' missing from evaluated snapshot")
                continue
            eval_ws = evaluated_wb[ws.title]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type != 'f':
                        continue
                    total_formulas += 1
                    try:
                        value_map.setdefault(ws.title, {})[cell.coordinate] = _coerce_cached_value(eval_ws[cell.coordinate].value, original_wb.epoch)
                    except Exception as exc:
                        warnings.append(f"Could not capture cached value for {ws.title}!{cell.coordinate}: {exc}")
    finally:
        original_wb.close(); evaluated_wb.close()
    return value_map, total_formulas, warnings


def _remove_calc_chain(parts: dict[str, bytes]) -> None:
    if 'xl/calcChain.xml' not in parts:
        return
    parts.pop('xl/calcChain.xml', None)
    if 'xl/_rels/workbook.xml.rels' in parts:
        root = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        changed = False
        for rel in list(root.findall('.//pr:Relationship', NS)):
            rel_type = rel.get('Type', '')
            target = rel.get('Target', '')
            if 'calcChain' in rel_type or target.endswith('calcChain.xml'):
                rel.getparent().remove(rel)
                changed = True
        if changed:
            parts['xl/_rels/workbook.xml.rels'] = ET.tostring(root, encoding='UTF-8', xml_declaration=True)
    if '[Content_Types].xml' in parts:
        root = ET.fromstring(parts['[Content_Types].xml'])
        changed = False
        for override in list(root.findall('.//ct:Override', NS)):
            if override.get('PartName') == '/xl/calcChain.xml':
                root.remove(override)
                changed = True
        if changed:
            parts['[Content_Types].xml'] = ET.tostring(root, encoding='UTF-8', xml_declaration=True)


def _ensure_calc_props(parts: dict[str, bytes]) -> None:
    if 'xl/workbook.xml' not in parts:
        return
    root = ET.fromstring(parts['xl/workbook.xml'])
    calc_pr = root.find('main:calcPr', NS)
    if calc_pr is None:
        calc_pr = ET.SubElement(root, f'{{{OOXML_MAIN_NS}}}calcPr')
    calc_pr.set('calcMode', 'auto')
    calc_pr.set('fullCalcOnLoad', '1')
    calc_pr.set('forceFullCalc', '1')
    calc_pr.set('calcOnSave', '1')
    parts['xl/workbook.xml'] = ET.tostring(root, encoding='UTF-8', xml_declaration=True)


def _update_sheet_xml(xml_bytes: bytes, values: dict[str, CachedValue]) -> tuple[bytes, int]:
    root = ET.fromstring(xml_bytes)
    count = 0
    for cell in root.findall('.//main:c[main:f]', NS):
        coord = cell.get('r')
        if not coord or coord not in values:
            continue
        cached = values[coord]
        if cached.xml_type:
            cell.set('t', cached.xml_type)
        else:
            cell.attrib.pop('t', None)
        v_node = cell.find('main:v', NS)
        if v_node is None:
            v_node = ET.SubElement(cell, f'{{{OOXML_MAIN_NS}}}v')
        v_node.text = cached.xml_value
        count += 1
    return ET.tostring(root, encoding='UTF-8', xml_declaration=True), count


def _inject_cached_values(workbook_path: Path, value_map: dict[str, dict[str, CachedValue]]) -> int:
    sheet_parts = {entry['name']: entry['part'] for entry in sheet_part_map(workbook_path) if entry.get('part')}
    with zipfile.ZipFile(workbook_path, 'r') as zf:
        parts = {name: zf.read(name) for name in zf.namelist()}
        compress_type = {info.filename: info.compress_type for info in zf.infolist()}
    updated = 0
    for sheet_name, coord_map in value_map.items():
        part = sheet_parts.get(sheet_name)
        if not part or part not in parts:
            continue
        parts[part], count = _update_sheet_xml(parts[part], coord_map)
        updated += count
    _remove_calc_chain(parts)
    _ensure_calc_props(parts)
    temp_output = workbook_path.with_suffix(workbook_path.suffix + '.tmp')
    with zipfile.ZipFile(temp_output, 'w') as zf:
        for name, data in parts.items():
            zf.writestr(name, data, compress_type=compress_type.get(name, zipfile.ZIP_DEFLATED))
    os.replace(temp_output, workbook_path)
    return updated


def _scan_cached_errors(workbook_path: Path) -> tuple[int, dict[str, dict[str, Any]]]:
    wb = load_workbook(workbook_path, data_only=True, keep_links=True)
    error_details = {err: [] for err in sorted(ERROR_LITERALS)}
    total_errors = 0
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in ERROR_LITERALS:
                        error_details[cell.value].append(f'{sheet_name}!{cell.coordinate}')
                        total_errors += 1
    finally:
        wb.close()
    return total_errors, {err: {'count': len(locs), 'locations': locs[:50]} for err, locs in error_details.items() if locs}


def recalc(filename: str | Path, timeout: int = 60) -> dict[str, Any]:
    path = Path(filename)
    if not path.exists():
        return {'error': f'File {path} does not exist'}
    if path.suffix.lower() not in EXCEL_PACKAGE_EXTENSIONS:
        return {'error': f'Unsupported Excel OOXML file type: {path.suffix.lower()}'}
    if is_probably_encrypted_excel(path):
        return {'error': 'Workbook appears to be password-protected or encrypted. Decrypt it first, or use msoffcrypto-tool if available.'}
    if not is_zip_package(path):
        return {'error': 'Workbook is not a valid OOXML ZIP package'}
    snapshot_path: Path | None = None
    try:
        snapshot_path = _run_libreoffice_snapshot(path, timeout=timeout)
        value_map, total_formulas, warnings = _formula_cell_map(path, snapshot_path)
        updated_formulas = _inject_cached_values(path, value_map)
        total_errors, error_summary = _scan_cached_errors(path)
        if updated_formulas < total_formulas:
            warnings.append(f'Only {updated_formulas} of {total_formulas} formula cells received cached values')
        return {
            'status': 'success' if total_errors == 0 else 'errors_found',
            'total_formulas': total_formulas,
            'cached_values_updated': updated_formulas,
            'total_errors': total_errors,
            'error_summary': error_summary,
            'warnings': warnings,
            'evaluation_backend': 'libreoffice_roundtrip_snapshot',
            'calc_chain_removed': True,
        }
    except subprocess.TimeoutExpired:
        return {'error': f'Timed out after {timeout}s while evaluating workbook'}
    except Exception as exc:
        return {'error': str(exc)}
    finally:
        if snapshot_path and snapshot_path.exists():
            snapshot_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description='Recalculate spreadsheet formulas and cache results back into the original OOXML workbook')
    parser.add_argument('excel_file', help='Path to .xlsx/.xlsm/.xltx/.xltm workbook')
    parser.add_argument('timeout', nargs='?', type=int, default=60, help='Timeout in seconds (default: 60)')
    args = parser.parse_args()
    print(json.dumps(recalc(args.excel_file, args.timeout), indent=2))

if __name__ == '__main__':
    main()
