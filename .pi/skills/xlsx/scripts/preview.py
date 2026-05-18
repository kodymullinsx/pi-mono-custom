"""Render an Excel workbook to per-sheet PNG previews.

Primary path:  xlsx2html (per sheet → HTML) → Playwright Chromium (full-page screenshot)
               One PNG per sheet, full content width, no page breaks.

Fallback path: LibreOffice → PDF → pdftoppm/pdftocairo rasterization
               Used automatically when xlsx2html or Playwright are unavailable or fail.

CLI usage:
    python scripts/preview.py workbook.xlsx [--output-dir DIR] [--dpi 144] [--timeout 60]
    python scripts/preview.py workbook.xlsx --method html          # force HTML path
    python scripts/preview.py workbook.xlsx --method libreoffice   # force legacy path
"""

from __future__ import annotations

import argparse
import io
import json
import shutil
import subprocess
import tempfile
import textwrap
from pathlib import Path
from typing import Any

from office.helpers.xlsx_package import EXCEL_PACKAGE_EXTENSIONS, is_probably_encrypted_excel, is_zip_package


# ---------------------------------------------------------------------------
# Capability probes (done once at import time)
# ---------------------------------------------------------------------------

def _has_module(name: str) -> bool:
    import importlib.util
    return importlib.util.find_spec(name) is not None

_HAS_XLSX2HTML  = _has_module("xlsx2html")
_HAS_PLAYWRIGHT = _has_module("playwright")


# ---------------------------------------------------------------------------
# HTML + Playwright path
# ---------------------------------------------------------------------------

def _sheet_names(workbook_path: Path) -> list[str]:
    """Return ordered sheet names from the workbook."""
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def _sheet_to_html(workbook_path: Path, sheet: str | None) -> str:
    """Convert one sheet to an HTML string via xlsx2html."""
    from xlsx2html import xlsx2html

    buf = io.StringIO()
    xlsx2html(str(workbook_path), buf, sheet=sheet, locale="en")
    raw_table = buf.getvalue()

    # Wrap in a minimal document that resets margin and sets a readable font.
    return textwrap.dedent(f"""\
        <!DOCTYPE html>
        <html>
        <head>
        <meta charset="UTF-8">
        <style>
          * {{ box-sizing: border-box; }}
          html, body {{
            margin: 0;
            padding: 12px;
            background: #ffffff;
            font-family: Calibri, Arial, sans-serif;
            font-size: 10pt;
          }}
          table {{
            border-collapse: collapse;
            white-space: pre-wrap;
            word-break: break-word;
          }}
          td, th {{
            overflow: hidden;
          }}
        </style>
        </head>
        <body>
        {raw_table}
        </body>
        </html>
    """)


def _playwright_screenshot(html: str, output_png: Path, min_width: int = 1200) -> None:
    """Render html string to a full-page PNG via headless Chromium."""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox"])
        # Start wide so tables don't collapse; we'll expand after content loads.
        page = browser.new_page(viewport={"width": min_width, "height": 900})
        page.set_content(html, wait_until="domcontentloaded")

        # Measure actual rendered content dimensions.
        dims = page.evaluate(
            "() => ({"
            "  w: Math.max(document.body.scrollWidth, document.documentElement.scrollWidth),"
            "  h: Math.max(document.body.scrollHeight, document.documentElement.scrollHeight)"
            "})"
        )
        final_w = max(dims["w"] + 24, min_width)
        final_h = max(dims["h"] + 24, 200)

        page.set_viewport_size({"width": final_w, "height": final_h})
        output_png.parent.mkdir(parents=True, exist_ok=True)
        page.screenshot(path=str(output_png), full_page=True, type="png")
        browser.close()


def _preview_html(workbook_path: Path, output_dir: Path, stem: str) -> tuple[list[str], list[str]]:
    """
    Generate one PNG per sheet using xlsx2html + Playwright.
    Returns (png_paths, warnings).
    """
    warnings: list[str] = []
    png_paths: list[str] = []

    try:
        sheets = _sheet_names(workbook_path)
    except Exception as exc:
        return [], [f"Could not read sheet names: {exc}"]

    for sheet in sheets:
        # Build a filesystem-safe filename from the sheet name.
        safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in sheet).strip()
        out_png = output_dir / f"{stem} - {safe}.png"

        try:
            html = _sheet_to_html(workbook_path, sheet)
        except Exception as exc:
            warnings.append(f"xlsx2html failed for sheet '{sheet}': {exc}")
            continue

        try:
            _playwright_screenshot(html, out_png)
            png_paths.append(str(out_png.resolve()))
        except Exception as exc:
            warnings.append(f"Screenshot failed for sheet '{sheet}': {exc}")

    return png_paths, warnings


# ---------------------------------------------------------------------------
# LibreOffice fallback path (original implementation)
# ---------------------------------------------------------------------------

def _rasterize_pdf(pdf_path: Path, output_dir: Path, stem: str, dpi: int) -> tuple[list[str], list[str]]:
    warnings: list[str] = []
    png_paths: list[str] = []
    prefix = output_dir / stem

    pdftoppm = shutil.which("pdftoppm")
    if pdftoppm:
        cmd = [pdftoppm, "-png", "-r", str(dpi), str(pdf_path), str(prefix)]
    else:
        pdftocairo = shutil.which("pdftocairo")
        if not pdftocairo:
            return [], ["Neither pdftoppm nor pdftocairo is available; returning PDF only."]
        cmd = [pdftocairo, "-png", "-r", str(dpi), str(pdf_path), str(prefix)]

    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        warnings.append(proc.stderr.strip() or proc.stdout.strip() or "PDF rasterization failed")
        return [], warnings

    png_paths = sorted(
        str(path.resolve())
        for path in output_dir.glob(f"{stem}*.png")
        if " - " not in path.name   # exclude html-path outputs
    )
    if not png_paths:
        warnings.append("Rasterizer completed but no PNG previews were produced.")
    return png_paths, warnings


def _preview_libreoffice(
    workbook_path: Path, output_dir: Path, stem: str, dpi: int, timeout: int
) -> tuple[str | None, list[str], list[str]]:
    """
    Convert workbook to PDF via LibreOffice, then rasterize to PNGs.
    Returns (pdf_path_str | None, png_paths, warnings).
    """
    from office.soffice import initialize_profile, run_soffice, temporary_profile

    warnings: list[str] = []

    with temporary_profile() as profile:
        initialize_profile(profile, timeout=min(timeout, 20))
        with tempfile.TemporaryDirectory(prefix="xlsx-preview-") as td:
            temp_dir = Path(td)
            source_copy = temp_dir / f"preview-source{workbook_path.suffix.lower()}"
            shutil.copy2(workbook_path, source_copy)

            proc = run_soffice(
                [
                    "--headless",
                    "--norestore",
                    "--nodefault",
                    "--nolockcheck",
                    "--nofirststartwizard",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_dir),
                    str(source_copy),
                ],
                profile_dir=profile,
                timeout=timeout,
                capture_output=True,
                text=True,
                check=False,
            )
            pdf_path = output_dir / f"{source_copy.stem}.pdf"
            if proc.returncode != 0 or not pdf_path.exists():
                detail = proc.stderr.strip() or proc.stdout.strip() or "unknown LibreOffice error"
                return None, [], [f"LibreOffice PDF export failed: {detail}"]

            final_pdf = output_dir / f"{stem}.pdf"
            if pdf_path != final_pdf:
                shutil.move(pdf_path, final_pdf)

            png_paths, raster_warnings = _rasterize_pdf(final_pdf, output_dir, stem, dpi)
            warnings.extend(raster_warnings)
            return str(final_pdf.resolve()), png_paths, warnings


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def preview_workbook(
    workbook: str | Path,
    output_dir: str | Path | None = None,
    dpi: int = 144,
    timeout: int = 60,
    method: str = "auto",   # "auto" | "html" | "libreoffice"
) -> dict[str, Any]:
    """
    Render a workbook to PNG previews.

    method="auto"  → xlsx2html+Playwright if available, else LibreOffice fallback
    method="html"  → force xlsx2html+Playwright (error if unavailable)
    method="libreoffice" → force legacy LibreOffice PDF path
    """
    path = Path(workbook)
    if not path.exists():
        return {"error": f"Workbook does not exist: {path}"}
    if path.suffix.lower() not in EXCEL_PACKAGE_EXTENSIONS:
        return {"error": f"Unsupported workbook extension '{path.suffix}'."}
    if is_probably_encrypted_excel(path):
        return {"error": "Workbook appears encrypted/password-protected. Decrypt before previewing."}
    if not is_zip_package(path):
        return {"error": "Workbook is not a valid OOXML ZIP package."}

    output_path = Path(output_dir) if output_dir else path.with_suffix("")
    output_path.mkdir(parents=True, exist_ok=True)
    stem = path.stem

    # --- Choose method -------------------------------------------------------
    use_html = False
    if method == "html":
        if not _HAS_XLSX2HTML or not _HAS_PLAYWRIGHT:
            missing = [m for m, ok in [("xlsx2html", _HAS_XLSX2HTML), ("playwright", _HAS_PLAYWRIGHT)] if not ok]
            return {"error": f"method='html' requested but missing: {', '.join(missing)}"}
        use_html = True
    elif method == "libreoffice":
        use_html = False
    else:  # auto
        use_html = _HAS_XLSX2HTML and _HAS_PLAYWRIGHT

    # --- HTML path -----------------------------------------------------------
    if use_html:
        png_paths, warnings = _preview_html(path, output_path, stem)
        if png_paths:
            return {
                "status": "success" if not warnings else "success_with_warnings",
                "method": "html",
                "workbook": str(path.resolve()),
                "pdf": None,
                "png_previews": png_paths,
                "page_count": len(png_paths),
                "warnings": warnings,
            }
        # If html path produced nothing (e.g. all sheets failed), fall through
        # unless the caller forced it.
        if method == "html":
            return {
                "error": "HTML preview produced no output.",
                "warnings": warnings,
            }
        # Otherwise fall through to LibreOffice.
        warnings_html = warnings
    else:
        warnings_html = []

    # --- LibreOffice path ----------------------------------------------------
    pdf_str, png_paths, lo_warnings = _preview_libreoffice(path, output_path, stem, dpi, timeout)
    all_warnings = warnings_html + lo_warnings

    if pdf_str is None and not png_paths:
        return {"error": "Both HTML and LibreOffice preview paths failed.", "warnings": all_warnings}

    return {
        "status": "success" if not all_warnings else "success_with_warnings",
        "method": "libreoffice",
        "workbook": str(path.resolve()),
        "pdf": pdf_str,
        "png_previews": png_paths,
        "page_count": len(png_paths) if png_paths else None,
        "warnings": all_warnings,
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Render an Excel workbook to PNG previews (one per sheet via xlsx2html+Playwright, "
                    "or per-page via LibreOffice fallback)."
    )
    parser.add_argument("workbook", help="Path to .xlsx/.xlsm/.xltx/.xltm workbook")
    parser.add_argument("--output-dir", help="Directory for output files")
    parser.add_argument("--dpi", type=int, default=144, help="PNG DPI for LibreOffice path (default: 144)")
    parser.add_argument("--timeout", type=int, default=60, help="LibreOffice timeout in seconds (default: 60)")
    parser.add_argument(
        "--method",
        choices=["auto", "html", "libreoffice"],
        default="auto",
        help="Preview method: auto (default), html (xlsx2html+Playwright), or libreoffice",
    )
    args = parser.parse_args()
    print(json.dumps(
        preview_workbook(args.workbook, args.output_dir, args.dpi, args.timeout, args.method),
        indent=2,
    ))


if __name__ == "__main__":
    main()
